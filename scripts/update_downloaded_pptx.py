import sys
import pptx
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

sys.stdout.reconfigure(encoding='utf-8')

pptx_path = r'C:\Users\Deepak\Downloads\EcoCast AI_ Executive Pitch Deck (Final 12 Slides).pptx'
prs = Presentation(pptx_path)

def update_text(text):
    original = text
    # Currency replacements
    text = text.replace('€79.68/t CBAM Shield', '₹7,450/t Tariff Shield')
    text = text.replace('€79.68/t', '₹7,450/t')
    text = text.replace('€79.68/tCO2 (~₹7,450 per tonne)', '₹7,450 per tonne (~€79.68/tCO2)')
    text = text.replace('€79.68/tCO2', '₹7,450/tCO2')
    text = text.replace('€80/t', '₹7,450/t')
    text = text.replace('€173/t steel', '₹16,200/t steel')
    text = text.replace('€', '₹')

    # Team & Roll number replacements
    text = text.replace('Roll: 23CH10028, ', '')
    text = text.replace('(Roll: 23CH10028) ', '')
    text = text.replace('Roll: 23CH10028', '')
    text = text.replace('Roll No: 23CH10028', '')
    text = text.replace('(Roll: 23CH10028)', '')

    # Deepak college addition
    if '• Deepak R (Team Lead - Systems Architect & Edge Developer)' in text and 'IIPE' not in text:
        text = text.replace(
            '• Deepak R (Team Lead - Systems Architect & Edge Developer)',
            '• Deepak R (Team Lead - Systems Architect & Edge Developer, Indian Institute of Petroleum and Energy - IIPE)'
        )
    if text.startswith('Deepak R (Team Lead):') and 'IIPE' not in text:
        text = text.replace(
            'Deepak R (Team Lead):',
            'Deepak R (Team Lead): Indian Institute of Petroleum and Energy (IIPE) |'
        )

    return text

def process_shape(shape):
    if shape.has_text_frame:
        for p in shape.text_frame.paragraphs:
            old_text = p.text
            new_text = update_text(old_text)
            if old_text != new_text:
                p.text = new_text
                print(f"Updated: '{old_text}' -> '{new_text}'")
    elif shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for sub in shape.shapes:
            process_shape(sub)

for idx, slide in enumerate(prs.slides):
    for shape in slide.shapes:
        process_shape(shape)

# Save both in-place and as a designated new copy
prs.save(pptx_path)
out_copy = r'C:\Users\Deepak\Downloads\EcoCast AI_ Executive Pitch Deck (Indian Rupees & IIPE).pptx'
prs.save(out_copy)
print(f"Saved updated presentation to {pptx_path} and {out_copy}")
