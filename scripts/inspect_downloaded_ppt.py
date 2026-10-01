import sys
import pptx
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

sys.stdout.reconfigure(encoding='utf-8')

pptx_path = r'C:\Users\Deepak\Downloads\EcoCast AI_ Executive Pitch Deck (Final with Links & Balanced Slide 10).pptx'
prs = Presentation(pptx_path)

def extract_texts(shape):
    texts = []
    if shape.has_text_frame:
        for p in shape.text_frame.paragraphs:
            t = p.text.strip()
            if t:
                texts.append(t)
    elif shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for sub in shape.shapes:
            texts.extend(extract_texts(sub))
    elif shape.has_table:
        for row in shape.table.rows:
            row_txt = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_txt:
                texts.append(" | ".join(row_txt))
    return texts

for i, slide in enumerate(prs.slides):
    print(f"\n==========================================")
    print(f"SLIDE {i+1}")
    print(f"==========================================")
    for shape in slide.shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            print(f"  [PICTURE: {shape.name}]")
        texts = extract_texts(shape)
        for t in texts:
            print(f"  • {t}")
